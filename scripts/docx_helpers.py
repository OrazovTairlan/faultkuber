import copy
import os

import docx
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.table import Table


class Report:
    def __init__(self, template, root):
        self.doc = docx.Document(template)
        self.root = root

    def save(self, path):
        self.doc.save(path)

    def table(self, first_header, second_header=None):
        for t in self.doc.tables:
            cells = t.rows[0].cells
            if cells[0].text.strip() == first_header and (second_header is None or cells[1].text.strip() == second_header):
                return t
        raise KeyError(first_header)

    def par(self, startswith):
        for p in self.doc.paragraphs:
            if p.text.strip().startswith(startswith):
                return p
        raise KeyError(startswith)

    def set_cell(self, cell, text, size=8.5, bold=False):
        paras = text if isinstance(text, list) else [text]
        first = cell.paragraphs[0]
        for extra in cell.paragraphs[1:]:
            extra._p.getparent().remove(extra._p)
        for r in list(first.runs):
            r._r.getparent().remove(r._r)
        for i, t in enumerate(paras):
            para = first if i == 0 else cell.add_paragraph()
            run = para.add_run(t)
            run.font.size = Pt(size)
            run.bold = bold
            para.paragraph_format.space_after = Pt(2)

    def fill(self, table, rows, size=8.5, start=1, first_col=0):
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                self.set_cell(table.rows[start + i].cells[first_col + j], val, size=size)

    def add_rows(self, table, rows, size=8.5):
        for row in rows:
            table._tbl.append(copy.deepcopy(table.rows[-1]._tr))
            for j, val in enumerate(row):
                self.set_cell(table.rows[-1].cells[j], val, size=size)

    def cursor(self, element):
        return Cursor(self, element)


class Cursor:
    def __init__(self, report, element):
        self.r = report
        self.doc = report.doc
        self.el = element._tbl if isinstance(element, Table) else element._p

    def _place(self, new_p):
        self.el.addnext(new_p._p)
        self.el = new_p._p
        return new_p

    def text(self, t, size=9.5, bold=False, italic=False, color=None, space_after=4):
        para = self.doc.add_paragraph()
        run = para.add_run(t)
        run.font.size, run.bold, run.italic = Pt(size), bold, italic
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        para.paragraph_format.space_after = Pt(space_after)
        return self._place(para)

    def heading(self, t, size=10.5):
        return self.text(t, size=size, bold=True, space_after=3)

    def bullets(self, items, size=9.5):
        for it in items:
            para = self.doc.add_paragraph(style="List Bullet")
            run = para.add_run(it)
            run.font.size = Pt(size)
            para.paragraph_format.space_after = Pt(1)
            self._place(para)

    def picture(self, rel, width=6.8, caption=None):
        para = self.doc.add_paragraph()
        para.add_run().add_picture(os.path.join(self.r.root, rel), width=Inches(width))
        self._place(para)
        if caption:
            self.text(caption, size=8.5, italic=True, color="52514E")

    def code(self, source, size=7):
        para = self.doc.add_paragraph()
        lines = source.rstrip().split("\n")
        for i, line in enumerate(lines):
            run = para.add_run(line)
            run.font.name, run.font.size = "Consolas", Pt(size)
            run._r.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Consolas")
            if i < len(lines) - 1:
                run.add_break(WD_BREAK.LINE)
        shade = OxmlElement("w:shd")
        shade.set(qn("w:val"), "clear")
        shade.set(qn("w:fill"), "F3F3F1")
        para._p.get_or_add_pPr().append(shade)
        para.paragraph_format.space_after = Pt(6)
        self._place(para)

    def table(self, rows, widths=None, size=8):
        t = self.doc.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                self.r.set_cell(t.rows[i].cells[j], val, size=size, bold=(i == 0))
        if widths:
            for row in t.rows:
                for j, w in enumerate(widths):
                    row.cells[j].width = Inches(w)
        self.el.addnext(t._tbl)
        self.el = t._tbl
        self.text("", size=4, space_after=2)
